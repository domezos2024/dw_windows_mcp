"""Document creation: Markdown (stdlib only) and Office files (python-docx /
openpyxl / python-pptx). Writes go through the same fs sandbox as fs.* tools."""
from __future__ import annotations

from ..auth.rbac import ToolCategory
from ..db import User
from . import tool
from .filesystem import _check_allowed


@tool(
    name="documents.create_markdown",
    category=ToolCategory.FILESYSTEM_WRITE,
    description="Erstellt eine Markdown-Datei (.md) mit optionalem Titel und Abschnitten.",
)
def create_markdown(
    user: User, path: str, title: str | None = None, sections: list[dict] | None = None, content: str | None = None
) -> dict:
    p = _check_allowed(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    if content is not None:
        text = content
    else:
        parts = []
        if title:
            parts.append(f"# {title}\n")
        for section in sections or []:
            heading = section.get("heading", "")
            body = section.get("body", "")
            if heading:
                parts.append(f"## {heading}\n")
            parts.append(f"{body}\n")
        text = "\n".join(parts)

    p.write_text(text, encoding="utf-8")
    return {"path": str(p), "bytes_written": len(text.encode("utf-8"))}


@tool(
    name="documents.create_docx",
    category=ToolCategory.FILESYSTEM_WRITE,
    description="Erstellt ein Word-Dokument (.docx) mit Titel und Absätzen.",
)
def create_docx(user: User, path: str, title: str | None = None, paragraphs: list[str] | None = None) -> dict:
    import docx

    p = _check_allowed(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()
    if title:
        doc.add_heading(title, level=1)
    for para in paragraphs or []:
        doc.add_paragraph(para)
    doc.save(str(p))
    return {"path": str(p)}


@tool(
    name="documents.create_xlsx",
    category=ToolCategory.FILESYSTEM_WRITE,
    description="Erstellt eine Excel-Datei (.xlsx) aus einer Liste von Zeilen (Liste von Listen).",
)
def create_xlsx(user: User, path: str, rows: list[list], sheet_name: str = "Sheet1") -> dict:
    import openpyxl

    p = _check_allowed(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    wb.save(str(p))
    return {"path": str(p), "rows_written": len(rows)}


@tool(
    name="documents.create_pptx",
    category=ToolCategory.FILESYSTEM_WRITE,
    description="Erstellt eine PowerPoint-Datei (.pptx) aus einer Liste von Folien ({title, body}).",
)
def create_pptx(user: User, path: str, slides: list[dict]) -> dict:
    from pptx import Presentation

    p = _check_allowed(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    layout = prs.slide_layouts[1]  # title + content
    for slide_spec in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_spec.get("title", "")
        body = slide.placeholders[1]
        body.text = slide_spec.get("body", "")
    prs.save(str(p))
    return {"path": str(p), "slides_written": len(slides)}
