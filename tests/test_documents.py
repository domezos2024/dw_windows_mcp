"""documents.* — real file creation, always confined to the sandbox temp dir
set up in conftest.py, so nothing touches the real filesystem outside it."""
from __future__ import annotations

import os

from dw_windows_mcp.auth.users import create_user
from dw_windows_mcp.db import Role
from dw_windows_mcp.tools import invoke_tool


def _sandbox_path(name: str) -> str:
    return os.path.join(os.environ["DW_MCP_FS_ROOTS"], name)


def test_create_markdown_with_sections():
    admin = create_user("doc_admin", "Pass1234", Role.STANDARD)
    dest = _sandbox_path("note.md")
    invoke_tool(
        admin,
        "documents.create_markdown",
        {"path": dest, "title": "Mein Titel", "sections": [{"heading": "Abschnitt 1", "body": "Text hier"}]},
    )
    with open(dest, encoding="utf-8") as f:
        text = f.read()
    assert "# Mein Titel" in text
    assert "## Abschnitt 1" in text
    assert "Text hier" in text


def test_create_docx_roundtrip():
    import docx

    admin = create_user("doc_admin2", "Pass1234", Role.STANDARD)
    dest = _sandbox_path("report.docx")
    invoke_tool(admin, "documents.create_docx", {"path": dest, "title": "Bericht", "paragraphs": ["Zeile eins"]})

    doc = docx.Document(dest)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Bericht" in all_text
    assert "Zeile eins" in all_text


def test_create_xlsx_roundtrip():
    import openpyxl

    admin = create_user("doc_admin3", "Pass1234", Role.STANDARD)
    dest = _sandbox_path("data.xlsx")
    invoke_tool(admin, "documents.create_xlsx", {"path": dest, "rows": [["Name", "Wert"], ["A", 1]]})

    wb = openpyxl.load_workbook(dest)
    ws = wb.active
    assert ws["A1"].value == "Name"
    assert ws["B2"].value == 1


def test_create_pptx_roundtrip():
    from pptx import Presentation

    admin = create_user("doc_admin4", "Pass1234", Role.STANDARD)
    dest = _sandbox_path("deck.pptx")
    invoke_tool(
        admin,
        "documents.create_pptx",
        {"path": dest, "slides": [{"title": "Folie 1", "body": "Inhalt"}]},
    )

    prs = Presentation(dest)
    assert len(prs.slides) == 1
    assert prs.slides[0].shapes.title.text == "Folie 1"
